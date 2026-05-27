"""PPTX parser using python-pptx for slide decks."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import TYPE_CHECKING

from any2poster.parsers.base import BaseParser
from any2poster.models import (
    ParsedDocument,
    ExtractedSection,
    ExtractedFigure,
    ExtractedTable,
    SectionType,
)

if TYPE_CHECKING:
    from pptx.presentation import Presentation
    from pptx.slide import Slide
    from pptx.shapes.base import BaseShape


_AFFIL_KEYWORDS = re.compile(
    r"\b(university|institute|laboratory|lab|department|college|school|"
    r"centre|center|research|academy|hospital|mit|stanford|cmu|ucla|"
    r"berkeley|eth|oxford|cambridge)\b",
    re.IGNORECASE,
)
_EMAIL_RE = re.compile(r"\b\S+@\S+\b")


@dataclass
class _TextShape:
    text: str
    top: int
    left: int
    width: int
    height: int
    max_font: float
    is_title: bool
    has_bullets: bool


@dataclass
class _SlideInfo:
    slide_idx: int
    slide_type: str  # title | divider | content
    title: str
    content: str
    figure_ids: list[str]
    table_ids: list[str]


class PPTXParser(BaseParser):
    """Parse PowerPoint presentations into poster-ready sections."""

    @property
    def supported_extensions(self) -> list[str]:
        return [".pptx"]

    def parse(self, input_path: str, output_dir: Path) -> ParsedDocument:
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError(
                "python-pptx is required for PPTX parsing. "
                "Install with: pip install python-pptx"
            )

        input_path = Path(input_path)
        pres: Presentation = Presentation(str(input_path))
        output_dir.mkdir(parents=True, exist_ok=True)

        slide_width = pres.slide_width
        slide_height = pres.slide_height

        figures: list[ExtractedFigure] = []
        tables: list[ExtractedTable] = []
        slide_infos: list[_SlideInfo] = []

        title = ""
        authors = ""
        affiliation = ""

        fig_counter = 0
        table_counter = 0

        for idx, slide in enumerate(pres.slides):
            bg_rgb = self._slide_background_rgb(slide)
            text_shapes = self._extract_text_shapes(
                slide=slide,
                slide_width=slide_width,
                slide_height=slide_height,
                bg_rgb=bg_rgb,
            )

            title_shape = self._find_title_shape(slide, text_shapes)
            is_title = self._is_title_slide(idx, text_shapes, title_shape)
            is_divider = self._is_divider_slide(text_shapes, title_shape) if not is_title else False

            # Title slide extraction
            if is_title and not title:
                title, authors, affiliation = self._extract_title_slide_metadata(
                    slide, text_shapes, title_shape
                )
                slide_infos.append(
                    _SlideInfo(
                        slide_idx=idx,
                        slide_type="title",
                        title=title,
                        content="",
                        figure_ids=[],
                        table_ids=[],
                    )
                )
                continue

            # Extract images and charts
            slide_figures: list[str] = []
            slide_tables: list[str] = []

            fig_ids, fig_counter, new_figures = self._extract_images_from_slide(
                slide=slide,
                output_dir=output_dir,
                slide_idx=idx,
                start_index=fig_counter,
                min_px=150,
            )
            slide_figures.extend(fig_ids)
            figures.extend(new_figures)

            tab_ids, table_counter, new_tables = self._extract_charts_from_slide(
                slide=slide,
                slide_idx=idx,
                start_index=table_counter,
            )
            slide_tables.extend(tab_ids)
            tables.extend(new_tables)

            # Build slide content text
            content = self._build_slide_content(text_shapes, title_shape)
            content = self._append_notes(slide, content)

            slide_infos.append(
                _SlideInfo(
                    slide_idx=idx,
                    slide_type="divider" if is_divider else "content",
                    title=title_shape.text.strip() if title_shape else (text_shapes[0].text if text_shapes else f"Slide {idx+1}"),
                    content=content,
                    figure_ids=slide_figures,
                    table_ids=slide_tables,
                )
            )

        sections = self._build_sections_from_slides(slide_infos)

        # Build full text for analytics
        full_text_parts = [title, authors, f"Affiliation: {affiliation}" if affiliation else ""]
        full_text_parts += [s.content for s in sections if s.content]
        full_text = "\n\n".join([p for p in full_text_parts if p]).strip()

        return ParsedDocument(
            source_path=str(input_path),
            source_format="pptx",
            title=title,
            authors=authors,
            affiliation=affiliation,
            sections=sections,
            figures=figures,
            tables=tables,
            raw_text=full_text,
            total_words=self._count_words(full_text),
        )

    # ──────────────────────────────────────────────────────────────
    # Slide classification + text extraction
    # ──────────────────────────────────────────────────────────────

    def _extract_text_shapes(
        self,
        slide: "Slide",
        slide_width: int,
        slide_height: int,
        bg_rgb: tuple[int, int, int] | None,
    ) -> list[_TextShape]:
        text_shapes: list[_TextShape] = []

        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape.width < 100000 or shape.height < 100000:
                continue
            if shape.width <= 0 or shape.height <= 0:
                continue
            if not self._is_within_slide(shape, slide_width, slide_height):
                continue
            text = self._extract_text_from_shape(shape)
            if not text.strip():
                continue
            if self._is_hidden_text(shape, bg_rgb):
                continue

            max_font = self._shape_max_font_size(shape)
            has_bullets = self._shape_has_bullets(shape)
            is_title = self._is_title_placeholder(shape)

            text_shapes.append(
                _TextShape(
                    text=text.strip(),
                    top=int(shape.top),
                    left=int(shape.left),
                    width=int(shape.width),
                    height=int(shape.height),
                    max_font=max_font,
                    is_title=is_title,
                    has_bullets=has_bullets,
                )
            )

        text_shapes.sort(key=lambda s: (s.top, s.left))
        return text_shapes

    def _find_title_shape(
        self, slide: "Slide", text_shapes: list[_TextShape]
    ) -> _TextShape | None:
        for ts in text_shapes:
            if ts.is_title:
                return ts
        if text_shapes:
            # fallback: largest font near top
            return max(text_shapes, key=lambda s: (s.max_font, -s.top))
        return None

    def _is_title_slide(
        self,
        slide_idx: int,
        text_shapes: list[_TextShape],
        title_shape: _TextShape | None,
    ) -> bool:
        if slide_idx <= 1 and title_shape:
            if len(text_shapes) <= 4 and title_shape.max_font >= 28:
                return True
        return False

    def _is_divider_slide(
        self,
        text_shapes: list[_TextShape],
        title_shape: _TextShape | None,
    ) -> bool:
        if not text_shapes:
            return False
        if len(text_shapes) == 1:
            ts = text_shapes[0]
            if ts.max_font >= 28 and len(ts.text.split()) <= 12 and not ts.has_bullets:
                return True
        if len(text_shapes) == 2 and title_shape:
            if title_shape.max_font >= 28 and not title_shape.has_bullets:
                return True
        return False

    def _build_slide_content(
        self, text_shapes: list[_TextShape], title_shape: _TextShape | None
    ) -> str:
        lines: list[str] = []
        for ts in text_shapes:
            if title_shape and ts.text == title_shape.text:
                continue
            if ts.text:
                lines.append(ts.text)
        return "\n".join(lines).strip()

    def _append_notes(self, slide: "Slide", content: str) -> str:
        notes = ""
        try:
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text or ""
                if len(notes_text.split()) >= 30:
                    notes = notes_text.strip()
        except Exception:
            notes = ""
        if notes:
            return (content + "\n\n[Presenter Notes]\n" + notes).strip() if content else (
                "[Presenter Notes]\n" + notes
            )
        return content

    # ──────────────────────────────────────────────────────────────
    # Title slide metadata
    # ──────────────────────────────────────────────────────────────

    def _extract_title_slide_metadata(
        self,
        slide: "Slide",
        text_shapes: list[_TextShape],
        title_shape: _TextShape | None,
    ) -> tuple[str, str, str]:
        title = title_shape.text.strip() if title_shape else (text_shapes[0].text if text_shapes else "Untitled")

        candidates = [ts.text for ts in text_shapes if not title_shape or ts.text != title_shape.text]
        authors = ""
        affiliation = ""

        # authors: email block or name-like block
        for text in candidates:
            if _EMAIL_RE.search(text):
                authors = text.strip()
                break

        if not authors:
            for text in candidates:
                if self._looks_like_author_block(text):
                    authors = text.strip()
                    break

        # affiliation: keyword match
        for text in candidates:
            if _AFFIL_KEYWORDS.search(text):
                affiliation = text.strip()
                break

        return title, authors, affiliation

    def _looks_like_author_block(self, text: str) -> bool:
        if len(text.split()) > 20:
            return False
        # Heuristic: multiple capitalized name-like tokens
        tokens = [t for t in re.split(r"[,;]|and", text) if t.strip()]
        name_like = 0
        for t in tokens:
            words = t.strip().split()
            if len(words) >= 2 and all(w[:1].isupper() for w in words if w):
                name_like += 1
        return name_like >= 1

    # ──────────────────────────────────────────────────────────────
    # Section construction
    # ──────────────────────────────────────────────────────────────

    def _build_sections_from_slides(self, slides: list[_SlideInfo]) -> list[ExtractedSection]:
        sections: list[ExtractedSection] = []
        section_id = 0

        has_dividers = any(s.slide_type == "divider" for s in slides)

        if not has_dividers:
            for s in slides:
                if s.slide_type != "content":
                    continue
                section_id += 1
                title = s.title or f"Slide {s.slide_idx + 1}"
                content = s.content or ""
                sections.append(
                    ExtractedSection(
                        section_id=f"section_{section_id}",
                        title=title,
                        section_type=SectionType(self._detect_section_type(title)),
                        content=content,
                        level=1,
                        word_count=self._count_words(content),
                        figure_ids=list(dict.fromkeys(s.figure_ids)),
                        table_ids=list(dict.fromkeys(s.table_ids)),
                    )
                )
            return sections

        current_title = ""
        current_content: list[str] = []
        current_figs: list[str] = []
        current_tables: list[str] = []

        def flush_section() -> None:
            nonlocal section_id, current_title, current_content, current_figs, current_tables
            if not current_title and not current_content:
                return
            section_id += 1
            content = "\n\n".join([c for c in current_content if c]).strip()
            sections.append(
                ExtractedSection(
                    section_id=f"section_{section_id}",
                    title=current_title or f"Section {section_id}",
                    section_type=SectionType(self._detect_section_type(current_title)),
                    content=content,
                    level=1,
                    word_count=self._count_words(content),
                    figure_ids=list(dict.fromkeys(current_figs)),
                    table_ids=list(dict.fromkeys(current_tables)),
                )
            )
            current_title = ""
            current_content = []
            current_figs = []
            current_tables = []

        for s in slides:
            if s.slide_type == "divider":
                flush_section()
                current_title = s.title or f"Section {section_id + 1}"
                continue
            if s.slide_type != "content":
                continue

            # If no divider seen yet, create an "Overview" section
            if not current_title:
                current_title = "Overview"

            if s.title and s.title.strip() and s.title.strip() != current_title:
                current_content.append(f"## {s.title.strip()}\n{s.content}".strip())
            else:
                current_content.append(s.content.strip())

            current_figs.extend(s.figure_ids)
            current_tables.extend(s.table_ids)

        flush_section()
        return sections

    # ──────────────────────────────────────────────────────────────
    # Images + charts
    # ──────────────────────────────────────────────────────────────

    def _extract_images_from_slide(
        self,
        slide: "Slide",
        output_dir: Path,
        slide_idx: int,
        start_index: int,
        min_px: int = 150,
    ) -> tuple[list[str], int, list[ExtractedFigure]]:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        figures: list[ExtractedFigure] = []
        figure_ids: list[str] = []
        counter = start_index

        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                image = shape.image
                if not image:
                    continue
                blob = image.blob
                if not blob:
                    continue
                # Determine extension safely — image.ext calls PIL internally
                # and will raise UnidentifiedImageError for EMF/WMF/SVG blobs
                try:
                    ext = image.ext or "png"
                except Exception:
                    ext = "png"
                counter += 1
                img_path = output_dir / f"figure_{counter}.{ext}"
                img_path.write_bytes(blob)

                if not self._is_image_large_enough(img_path, min_px=min_px):
                    try:
                        img_path.unlink()
                    except Exception:
                        pass
                    counter -= 1
                    continue

                caption = self._nearest_text_caption(slide, shape)
                fig_id = f"fig_{counter}"
                figures.append(
                    ExtractedFigure(
                        figure_id=fig_id,
                        caption=caption,
                        page_number=slide_idx + 1,
                        path=img_path,
                    )
                )
                figure_ids.append(fig_id)
            except Exception:
                continue

        return figure_ids, counter, figures

    def _extract_charts_from_slide(
        self,
        slide: "Slide",
        slide_idx: int,
        start_index: int,
    ) -> tuple[list[str], int, list[ExtractedTable]]:
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        tables: list[ExtractedTable] = []
        table_ids: list[str] = []
        counter = start_index

        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.CHART:
                continue
            chart = shape.chart
            counter += 1
            rows = self._chart_to_rows(chart)
            if not rows:
                continue
            content = self._format_table(rows)
            table_id = f"table_{counter}"
            chart_title = self._get_chart_title(chart)
            caption = chart_title if chart_title else f"Chart {counter} (slide {slide_idx + 1})"
            tables.append(
                ExtractedTable(
                    table_id=table_id,
                    caption=caption,
                    content=content,
                    page_number=slide_idx + 1,
                )
            )
            table_ids.append(table_id)

        return table_ids, counter, tables

    def _get_chart_title(self, chart) -> str:
        try:
            if chart.has_title and chart.chart_title.text_frame:
                return chart.chart_title.text_frame.text.strip()
        except Exception:
            pass
        return ""

    def _chart_to_rows(self, chart) -> list[list[str]]:
        try:
            plots = chart.plots
            if not plots:
                return []
            plot = plots[0]

            categories: list[str] = []
            if plot.categories:
                for c in plot.categories:
                    try:
                        label = c.label
                        categories.append(str(label) if label is not None else "")
                    except Exception:
                        categories.append("")

            series = list(plot.series)
            if not series:
                return []

            # Fall back to index-based categories for charts without a category axis
            if not categories:
                try:
                    n = len(list(series[0].values))
                    categories = [str(i + 1) for i in range(n)]
                except Exception:
                    return []

            headers = ["Category"] + [s.name or f"Series {i+1}" for i, s in enumerate(series)]
            rows = [headers]
            for idx, cat in enumerate(categories):
                row = [cat]
                for s in series:
                    try:
                        v = s.values[idx]
                        if v is None:
                            row.append("")
                        elif isinstance(v, float):
                            row.append(str(round(v, 4)))
                        else:
                            row.append(str(v))
                    except Exception:
                        row.append("")
                rows.append(row)
            return rows
        except Exception:
            return []

    # ──────────────────────────────────────────────────────────────
    # Helpers
    # ──────────────────────────────────────────────────────────────

    def _is_title_placeholder(self, shape: "BaseShape") -> bool:
        try:
            from pptx.enum.shapes import PP_PLACEHOLDER
            if shape.is_placeholder:
                ptype = shape.placeholder_format.type
                return ptype in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) or shape.placeholder_format.idx == 0
        except Exception:
            pass
        return False

    def _shape_has_bullets(self, shape: "BaseShape") -> bool:
        try:
            tf = shape.text_frame
            for p in tf.paragraphs:
                if p.level and p.level > 0:
                    return True
        except Exception:
            return False
        return False

    def _extract_text_from_shape(self, shape: "BaseShape") -> str:
        try:
            tf = shape.text_frame
        except Exception:
            return ""
        lines: list[str] = []
        for p in tf.paragraphs:
            text = (p.text or "").strip()
            if not text:
                continue
            prefix = "- " if p.level and p.level >= 1 else ""
            lines.append(prefix + text)
        return "\n".join(lines)

    def _shape_max_font_size(self, shape: "BaseShape") -> float:
        max_size = 0.0
        try:
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    if run.font.size is not None:
                        max_size = max(max_size, run.font.size.pt)
        except Exception:
            pass
        return max_size or 12.0

    def _slide_background_rgb(self, slide: "Slide") -> tuple[int, int, int] | None:
        try:
            fill = slide.background.fill
            if fill and fill.type is not None and fill.fore_color and fill.fore_color.rgb:
                rgb = fill.fore_color.rgb
                return (rgb[0], rgb[1], rgb[2])
        except Exception:
            pass
        return (255, 255, 255)

    def _is_hidden_text(self, shape: "BaseShape", bg_rgb: tuple[int, int, int] | None) -> bool:
        if not bg_rgb:
            return False
        try:
            saw_color = False
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    if run.font.color and run.font.color.rgb:
                        saw_color = True
                        rgb = run.font.color.rgb
                        dist = abs(rgb[0] - bg_rgb[0]) + abs(rgb[1] - bg_rgb[1]) + abs(rgb[2] - bg_rgb[2])
                        if dist > 30:
                            return False
            if not saw_color:
                return False
            return True
        except Exception:
            return False

    def _is_within_slide(self, shape: "BaseShape", slide_width: int, slide_height: int) -> bool:
        try:
            if shape.left < 0 or shape.top < 0:
                return False
            if shape.left + shape.width > slide_width + 1:
                return False
            if shape.top + shape.height > slide_height + 1:
                return False
        except Exception:
            return True
        return True

    def _nearest_text_caption(self, slide: "Slide", target_shape: "BaseShape") -> str:
        best = ""
        best_dist = None
        tx, ty = int(target_shape.left + target_shape.width // 2), int(target_shape.top + target_shape.height // 2)
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = (shape.text_frame.text or "").strip()
            if not text:
                continue
            cx, cy = int(shape.left + shape.width // 2), int(shape.top + shape.height // 2)
            dist = abs(cx - tx) + abs(cy - ty)
            if best_dist is None or dist < best_dist:
                best_dist = dist
                best = text
        return best

    def _is_image_large_enough(self, path: Path, min_px: int = 150) -> bool:
        try:
            from PIL import Image
            with Image.open(path) as img:
                w, h = img.size
            return w >= min_px and h >= min_px
        except Exception:
            return True

    def _format_table(self, rows: list[list[str]]) -> str:
        header = rows[0]
        body = rows[1:] if len(rows) > 1 else []
        header_line = "| " + " | ".join(header) + " |"
        sep_line = "| " + " | ".join(["---"] * len(header)) + " |"
        body_lines = ["| " + " | ".join(r) + " |" for r in body]
        return "\n".join([header_line, sep_line] + body_lines)
