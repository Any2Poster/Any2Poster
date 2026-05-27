"""
DOCX parser using python-docx for Word documents.
"""

from __future__ import annotations

from pathlib import Path
import re
from typing import TYPE_CHECKING, Iterable

from any2poster.parsers.base import BaseParser
from any2poster.models import (
    ParsedDocument,
    ExtractedSection,
    ExtractedFigure,
    ExtractedTable,
    SectionType,
)

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument
    from docx.text.paragraph import Paragraph
    from docx.table import Table


_CAPTION_RE = re.compile(r"^(?:Figure|Fig\.?)\s*\d+[\.:]\s*", re.IGNORECASE)
_SECTION_KEYWORDS = re.compile(
    r"^(abstract|introduction|background|related work|methods?|"
    r"methodology|approach|results?|discussion|conclusion|references)\b",
    re.IGNORECASE,
)


class DOCXParser(BaseParser):
    """Parse Word documents using python-docx."""

    @property
    def supported_extensions(self) -> list[str]:
        return [".docx", ".doc"]

    def parse(self, input_path: str, output_dir: Path) -> ParsedDocument:
        """Parse a DOCX file and extract content."""
        try:
            from docx import Document
        except ImportError:
            raise ImportError(
                "python-docx is required for DOCX parsing. "
                "Install with: pip install python-docx"
            )

        input_path = Path(input_path)
        if input_path.suffix.lower() == ".doc":
            raise ValueError(
                "Legacy .doc files are not supported. Please convert to .docx first."
            )

        doc = Document(str(input_path))

        paragraphs = list(doc.paragraphs)
        sections, heading_map, full_text = self._extract_sections(paragraphs)

        # Extract figures + tables and associate with sections
        figures, tables = self._extract_tables_and_figures(
            doc=doc,
            output_dir=output_dir,
            heading_map=heading_map,
            sections=sections,
        )

        # Title/authors
        title = self._extract_title(doc, paragraphs, sections)
        authors = self._extract_authors(doc, paragraphs, title)

        return ParsedDocument(
            source_path=str(input_path),
            source_format="docx",
            title=title,
            authors=authors,
            sections=sections,
            figures=figures,
            tables=tables,
            raw_text=full_text,
            total_words=self._count_words(full_text),
        )

    # ──────────────────────────────────────────────────────────────
    # Section detection
    # ──────────────────────────────────────────────────────────────

    def _extract_sections(
        self, paragraphs: list["Paragraph"]
    ) -> tuple[list[ExtractedSection], dict[object, str], str]:
        headings = self._detect_headings(paragraphs)

        sections: list[ExtractedSection] = []
        heading_map: dict[object, str] = {}
        current_section: ExtractedSection | None = None
        current_content: list[str] = []
        section_id = 0
        full_text_parts: list[str] = []

        for para in paragraphs:
            text = para.text.strip()
            if text:
                full_text_parts.append(text)

            if para._p in headings:
                # finalize previous section
                if current_section:
                    current_section.content = "\n".join(current_content).strip()
                    if current_section.content:
                        sections.append(current_section)

                section_id += 1
                level = headings[para._p]
                current_section = ExtractedSection(
                    section_id=f"section_{section_id}",
                    title=text or f"Section {section_id}",
                    content="",
                    section_type=SectionType(self._detect_section_type(text)),
                    level=level,
                )
                heading_map[para._p] = current_section.section_id
                current_content = []
            else:
                if text:
                    current_content.append(text)

        if current_section:
            current_section.content = "\n".join(current_content).strip()
            if current_section.content:
                sections.append(current_section)

        full_text = "\n".join(full_text_parts).strip()

        if not sections and full_text:
            sections.append(
                ExtractedSection(
                    section_id="section_1",
                    title="Content",
                    content=full_text,
                    section_type=SectionType.OTHER,
                    level=1,
                )
            )

        return sections, heading_map, full_text

    def _detect_headings(self, paragraphs: list["Paragraph"]) -> dict[object, int]:
        headings: dict[object, int] = {}

        # Pass 1: Word heading styles
        for para in paragraphs:
            if self._is_heading_style(para):
                level = self._heading_level_from_style(para)
                headings[para._p] = level

        if len(headings) >= 2:
            return headings

        # Pass 2: Heuristic headings based on formatting
        headings = {}
        body_size = self._estimate_body_font_size(paragraphs)
        candidates: list[tuple[object, float]] = []

        for para in paragraphs:
            if self._is_heading_heuristic(para, body_size):
                size = self._para_max_font_size(para) or body_size or 11.0
                candidates.append((para._p, size))

        if not candidates:
            return headings

        # Rank sizes to derive heading levels
        unique_sizes = sorted({s for _, s in candidates}, reverse=True)
        size_to_level = {s: idx + 1 for idx, s in enumerate(unique_sizes)}

        for p, size in candidates:
            headings[p] = size_to_level.get(size, 1)

        return headings

    def _is_heading_style(self, para: "Paragraph") -> bool:
        try:
            return para.style and para.style.name.startswith("Heading")
        except Exception:
            return False

    def _heading_level_from_style(self, para: "Paragraph") -> int:
        try:
            name = para.style.name
            m = re.search(r"Heading\s*(\d+)", name)
            if m:
                return int(m.group(1))
        except Exception:
            pass
        return 1

    def _is_heading_heuristic(self, para: "Paragraph", body_size: float | None) -> bool:
        text = para.text.strip()
        if not text:
            return False
        if self._is_bulleted(para):
            return False
        # short uppercase headings
        words = text.split()
        has_letters = any(c.isalpha() for c in text)
        if text.upper() == text and has_letters and len(words) <= 8:
            return True

        # bold + larger font size
        all_bold = self._para_all_bold(para)
        size = self._para_max_font_size(para)
        if all_bold and size and body_size and size >= body_size + 1:
            return True

        return False

    def _is_bulleted(self, para: "Paragraph") -> bool:
        try:
            ppr = para._p.pPr
            if ppr is not None and ppr.numPr is not None:
                return True
        except Exception:
            pass
        return False

    def _para_all_bold(self, para: "Paragraph") -> bool:
        runs = [r for r in para.runs if r.text and r.text.strip()]
        if not runs:
            return False
        para_bold = False
        try:
            para_bold = bool(para.style and para.style.font and para.style.font.bold)
        except Exception:
            para_bold = False
        for run in runs:
            if run.bold is True:
                continue
            if run.bold is None and para_bold:
                continue
            if run.font and run.font.bold:
                continue
            return False
        return True

    def _para_max_font_size(self, para: "Paragraph") -> float | None:
        sizes: list[float] = []
        for run in para.runs:
            if not run.text or not run.text.strip():
                continue
            size = None
            if run.font is not None and run.font.size is not None:
                size = run.font.size.pt
            elif run.style is not None and run.style.font is not None and run.style.font.size is not None:
                size = run.style.font.size.pt
            if size:
                sizes.append(size)
        if not sizes:
            try:
                if para.style and para.style.font and para.style.font.size:
                    return para.style.font.size.pt
            except Exception:
                return None
            return None
        return max(sizes)

    def _estimate_body_font_size(self, paragraphs: list["Paragraph"]) -> float | None:
        sizes: list[float] = []
        for para in paragraphs:
            text = para.text.strip()
            if not text or len(text.split()) < 5:
                continue
            if self._is_heading_style(para):
                continue
            if text.upper() == text and any(c.isalpha() for c in text):
                continue
            size = self._para_max_font_size(para)
            if size:
                sizes.append(size)
        if not sizes:
            return 11.0
        sizes.sort()
        return sizes[len(sizes) // 2]

    # ──────────────────────────────────────────────────────────────
    # Title / author extraction
    # ──────────────────────────────────────────────────────────────

    def _extract_title(
        self,
        doc: "DocxDocument",
        paragraphs: list["Paragraph"],
        sections: list[ExtractedSection],
    ) -> str:
        if doc.core_properties.title:
            return doc.core_properties.title

        for para in paragraphs[:10]:
            try:
                if para.style and para.style.name.lower() == "title" and para.text.strip():
                    return para.text.strip()
            except Exception:
                pass

        if sections:
            return sections[0].title

        for para in paragraphs:
            if para.text.strip():
                return para.text.strip()[:200]

        return "Untitled Document"

    def _extract_authors(
        self,
        doc: "DocxDocument",
        paragraphs: list["Paragraph"],
        title: str,
    ) -> str:
        # Core properties
        if doc.core_properties.author:
            return doc.core_properties.author.strip()

        # Style-based hints
        for para in paragraphs[:12]:
            name = ""
            try:
                name = (para.style.name or "").lower()
            except Exception:
                name = ""
            if name in {"author", "authors", "subtitle"} and para.text.strip():
                return para.text.strip()

        # Heuristic: paragraphs right after title
        title_idx = None
        for idx, para in enumerate(paragraphs[:20]):
            if para.text.strip() and para.text.strip() == title:
                title_idx = idx
                break

        if title_idx is None:
            return ""

        candidates: list[str] = []
        for para in paragraphs[title_idx + 1:title_idx + 4]:
            text = para.text.strip()
            if not text:
                continue
            if self._is_heading_style(para):
                continue
            if self._is_bulleted(para):
                continue
            if _SECTION_KEYWORDS.match(text):
                continue
            if len(text.split()) > 20:
                continue
            candidates.append(text)

        if candidates:
            return ", ".join(candidates)

        return ""

    # ──────────────────────────────────────────────────────────────
    # Figures + tables (with section association)
    # ──────────────────────────────────────────────────────────────

    def _extract_tables_and_figures(
        self,
        doc: "DocxDocument",
        output_dir: Path,
        heading_map: dict[object, str],
        sections: list[ExtractedSection],
    ) -> tuple[list[ExtractedFigure], list[ExtractedTable]]:
        figures: list[ExtractedFigure] = []
        tables: list[ExtractedTable] = []
        figures_dir = output_dir
        figures_dir.mkdir(exist_ok=True)

        section_map = {s.section_id: s for s in sections}
        current_section_id = sections[0].section_id if sections else ""

        blocks = list(self._iter_block_items(doc))
        fig_counter = 0
        table_counter = 0

        for idx, block in enumerate(blocks):
            if self._is_paragraph(block):
                para = block
                if para._p in heading_map:
                    current_section_id = heading_map[para._p]

                # Inline images in this paragraph
                for rel_id in self._iter_image_rel_ids(para):
                    image_part = doc.part.related_parts.get(rel_id)
                    if not image_part:
                        continue
                    fig_counter += 1
                    ext = self._image_ext_from_content_type(image_part.content_type)
                    img_path = figures_dir / f"figure_{fig_counter}{ext}"
                    img_path.write_bytes(image_part.blob)

                    if not self._is_image_large_enough(img_path, min_px=200):
                        try:
                            img_path.unlink()
                        except Exception:
                            pass
                        continue

                    caption = self._find_caption_after(blocks, idx)

                    figure_id = f"fig_{fig_counter}"
                    figures.append(
                        ExtractedFigure(
                            figure_id=figure_id,
                            path=img_path,
                            caption=caption,
                        )
                    )
                    if current_section_id and current_section_id in section_map:
                        section_map[current_section_id].figure_ids.append(figure_id)

            elif self._is_table(block):
                table_counter += 1
                rows = [
                    [cell.text.strip() for cell in row.cells]
                    for row in block.rows
                ]
                rows = [r for r in rows if any(c for c in r)]
                if not rows:
                    continue

                content = self._format_table(rows)
                table_id = f"table_{table_counter}"
                tables.append(
                    ExtractedTable(
                        table_id=table_id,
                        caption=f"Table {table_counter}",
                        content=content,
                    )
                )
                if current_section_id and current_section_id in section_map:
                    section_map[current_section_id].table_ids.append(table_id)

        return figures, tables

    def _iter_block_items(self, doc: "DocxDocument") -> Iterable[object]:
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in doc.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, doc)
            elif isinstance(child, CT_Tbl):
                yield Table(child, doc)

    def _is_paragraph(self, block: object) -> bool:
        return block.__class__.__name__ == "Paragraph"

    def _is_table(self, block: object) -> bool:
        return block.__class__.__name__ == "Table"

    def _iter_image_rel_ids(self, para: "Paragraph") -> Iterable[str]:
        try:
            from docx.oxml.ns import qn
        except Exception:
            return []
        rel_ids: list[str] = []
        for blip in para._p.xpath(".//a:blip"):
            rid = blip.get(qn("r:embed"))
            if rid:
                rel_ids.append(rid)
        return rel_ids

    def _find_caption_after(self, blocks: list[object], idx: int) -> str:
        for j in range(idx + 1, min(idx + 4, len(blocks))):
            b = blocks[j]
            if not self._is_paragraph(b):
                continue
            text = b.text.strip()
            if not text:
                continue
            if _CAPTION_RE.match(text):
                return text
            # Accept short caption-like lines
            if len(text.split()) <= 12 and text.lower().startswith("figure"):
                return text
            break
        return ""

    def _image_ext_from_content_type(self, content_type: str) -> str:
        ct = (content_type or "").lower()
        if "jpeg" in ct or "jpg" in ct:
            return ".jpg"
        if "gif" in ct:
            return ".gif"
        if "bmp" in ct:
            return ".bmp"
        if "tiff" in ct or "tif" in ct:
            return ".tif"
        return ".png"

    def _is_image_large_enough(self, path: Path, min_px: int = 200) -> bool:
        try:
            from PIL import Image
            with Image.open(path) as img:
                w, h = img.size
            return w >= min_px and h >= min_px
        except Exception:
            return True

    def _format_table(self, rows: list[list[str]]) -> str:
        header = rows[0]
        body = rows[1:] if len(rows) > 1 else []
        header_line = "| " + " | ".join(header) + " |"
        sep_line = "| " + " | ".join(["---"] * len(header)) + " |"
        body_lines = ["| " + " | ".join(r) + " |" for r in body]
        return "\n".join([header_line, sep_line] + body_lines)
